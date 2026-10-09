# -*- coding: utf-8 -*-
"""Alocação por classe de ativo da pessoa física por segmento (ANBIMA): Varejo consolidado, Varejo Tradicional, Varejo Alta Renda e Private,
em cada edição baixada (fontes/setor/anbima/historico/{varejo,private}_*.xls[x] + Estatistica_de_{Varejo,Private}_AAAAMM.xlsx).
A data de referência é lida DENTRO do arquivo ("Ago/26", "Dezembro-18"), não do nome (que é a data de publicação/captura).
Varejo: aba "Volume Financeiro", três blocos ("TOTAL", "VAREJO TRADICIONAL", "VAREJO ALTA RENDA"); rótulo numa coluna, TOTAL Brasil na seguinte;
edições de 2017-18 usam siglas (LCA, CRI...). Private layout antigo (2018-jan/24): rótulo na col B, DUAS colunas de valor (dezembro do ano
anterior na col C e o mês corrente na col E) → cada arquivo rende dois pontos; estruturados vêm dentro de "1. FUNDOS". Layout novo (mai/24+):
rótulo col C, valor col D, seções 1-6. Quando várias edições trazem a mesma referência, fica a mais recente (revisada).
Saídas: data/anbima_alocacao_segmentos_long.csv (seg, data AAAAMM, classe, valor R$ mi, arquivo) e data/anbima_alocacao_segmentos.csv (largo,
com agregados lci_lca, credito_privado, estruturados, fundos_555)."""
import os, re, glob, unicodedata, datetime as dt, pandas as pd
HERE = os.path.dirname(os.path.abspath(__file__)); AN = os.path.join(HERE, "fontes", "setor", "anbima"); DATA = os.path.join(HERE, "data")
MESES = {m: i + 1 for i, m in enumerate(["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"])}

def norm(s): return re.sub(r"\s+", " ", re.sub(r"\[\d\]", "", unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower())).strip()

def ref(tok):
    """'Ago/26', 'Jun/21**', 'Dezembro-18*', datetime → AAAAMM; None se não for data."""
    if hasattr(tok, "year"): return tok.year * 100 + tok.month
    m = re.match(r"^\s*([a-zç]{3})[a-zç]*\s*[/\-]\s*(\d{2,4})\b", norm(tok))
    if not m or m.group(1) not in MESES: return None
    y = int(m.group(2)); y = y + 2000 if y < 100 else y
    return y * 100 + MESES[m.group(1)]

def grade(path):
    """Linhas como listas de células (xlsx via openpyxl, xls via xlrd, ods via pandas/odf)."""
    if path.lower().endswith(".ods"):
        df = pd.read_excel(path, engine="odf", sheet_name=0, header=None); return [[None if (isinstance(v, float) and pd.isna(v)) else v for v in r] for r in df.values.tolist()]
    if path.lower().endswith(".xlsx"):
        import openpyxl; wb = openpyxl.load_workbook(path, read_only=True, data_only=True); return [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    import xlrd; ws = xlrd.open_workbook(path).sheet_by_index(0); return [ws.row_values(i) for i in range(ws.nrows)]

def num(v): return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

def casa(n, pref): return n == pref or n.startswith(pref + " ") or n.startswith(pref + "/") or n.startswith(pref + "(") or n == pref + "s" or n.startswith(pref + "s ") or (len(pref) > 6 and n.startswith(pref))

# ---------- VAREJO ----------
V_SECOES = {"1. fundos de investimento": "1", "2. fundos estruturados": "2", "3. titulos e valores": "3", "4. poupanca": "4", "5. previdencia": "5"}
V_SEC_CLASSE = {"1": "fundos_555", "2": "estruturados", "3": "tvm", "4": "poupanca", "5": "previdencia"}
V_MAPA = [("1", "renda fixa (baixa", "fundos_rf_baixa_dur"), ("1", "renda fixa (exceto", "fundos_rf_exceto_bd"), ("1", "renda fixa", "fundos_rf"), ("1", "multimercados", "fundos_mm"),
          ("1", "acoes", "fundos_acoes"), ("1", "fmp", "fmp"), ("1", "cambial", "fundos_cambial"), ("1", "outros", "fundos_outros"),
          ("2", "fundo de investimento em direitos", "fidc"), ("2", "fidc", "fidc"), ("2", "fundo de investimento imobiliario", "fii"), ("2", "fii", "fii"), ("2", "fundo de investimento em participacoes", "fip"), ("2", "fip", "fip"),
          ("2", "etf renda fixa", "etf_rf"), ("2", "etf renda variavel", "etf_rv"), ("2", "etf", "etf"),
          ("3", "acoes", "acoes"), ("3", "titulos publicos", "tit_publicos"), ("3", "pre-fixado", "tp_pre"), ("3", "pos-fixado", "tp_pos"), ("3", "hibrido", "tp_hibrido"), ("3", "titulos privados", "tit_privados"),
          ("3", "cdb", "cdb"), ("3", "op. compromissada", "compromissada"), ("3", "letras de credito agr", "lca"), ("3", "lca", "lca"), ("3", "certificado de recebiveis agr", "cra"), ("3", "cra", "cra"),
          ("3", "letras de credito imob", "lci"), ("3", "lci", "lci"), ("3", "certificado de recebiveis imob", "cri"), ("3", "cri", "cri"), ("3", "letras financeiras", "lf"), ("3", "lf", "lf"),
          ("3", "debentures tradicionais", "deb_tradicionais"), ("3", "debentures incentivadas", "deb_incentivadas"), ("3", "debentures", "debentures"), ("3", "letras hipotecarias", "lh"), ("3", "lh", "lh"),
          ("3", "letra imobiliaria garantida", "lig"), ("3", "letra de arrendamento", "lam"), ("3", "letra de cambio", "lc"), ("3", "box", "box"), ("3", "outros", "tvm_outros"),
          ("3", "certificado de operacoes estruturadas", "coe"), ("3", "coe", "coe"), ("3", "renda fixa", "tvm_rf"), ("3", "renda variavel", "tvm_rv")]
V_BLOCOS = [("varejo tradicional", "varejo_tradicional"), ("varejo alta renda", "varejo_alta_renda"), ("varejo", "varejo_tradicional")]  # 2014-15: bloco "VAREJO" = tradicional

def parse_varejo(path):
    G = grade(path); data = next((ref(v) for r in G[:10] for v in r if v is not None and ref(v)), None)
    rows = []; seg = None; sec = None; vistos = set()
    for row in G:
        j = next((j for j, v in enumerate(row) if isinstance(v, str) and v.strip()), None)
        if j is None: continue
        n = norm(row[j]); val = next((num(v) for v in row[j + 1:] if num(v) is not None), None)
        if seg is None:
            if n.startswith("total") and val is not None: seg = "varejo_total"; rows.append((seg, "total", val))
            continue
        hit = next((s for k, s in V_BLOCOS if n.startswith(k)), None)
        if hit and val is not None: seg = hit; sec = None; vistos = set(); rows.append((seg, "total", val)); continue
        sh = next((s for k, s in V_SECOES.items() if n.startswith(k)), None)
        if sh: sec = sh; rows.append((seg, V_SEC_CLASSE[sh], val)); continue
        if sec is None or val is None: continue
        for s, pref, cl in V_MAPA:
            if s == sec and casa(n, pref) and (seg, cl) not in vistos: rows.append((seg, cl, val)); vistos.add((seg, cl)); break
    return pd.DataFrame(rows, columns=["seg", "classe", "valor"]).assign(data=data, arquivo=os.path.basename(path))

# ---------- PRIVATE ----------
P_SECOES = [("i - volume financeiro", "total"), ("i - posicao de aum", "total"), ("total", "total"), ("1. fundos", "fundos"), ("fundos estruturados", "estruturados"), ("2. fundos estruturados", "estruturados"),
            ("2. titulos e valores", "tvm"), ("3. titulos e valores", "tvm"), ("3. caixa", "caixa_poupanca"), ("4. caixa", "caixa_poupanca"), ("3. caixa / poupanca", "caixa_poupanca"),
            ("4. previdencia", "previdencia"), ("5. previdencia", "previdencia"), ("5. outros investimentos", "outros_invest"), ("6. posicao de credito", "fim"), ("iii - posicao de credito", "fim"), ("ii - domicilio", "fim"), ("numero de contas", "fim")]
P_MAPA = {"fundos": [("fundos abertos proprios", "fundos_abertos_proprios"), ("fundos proprios", "fundos_abertos_proprios"), ("fundos abertos de terceiros", "fundos_abertos_terceiros"), ("fundos terceiros", "fundos_abertos_terceiros"), ("fundos abertos", "fundos_abertos"), ("fundos exclusivos", "fundos_exclusivos")],
          "estruturados": [("participacoes", "fip"), ("fundo de investimento em participacoes", "fip"), ("fidc", "fidc"), ("fundo de investimento em direitos", "fidc"), ("imobiliario", "fii"), ("fundo de investimento imobiliario", "fii"), ("etf", "etf"), ("fmp", "fmp"), ("outros", "estr_outros")],
          "tvm": [("acoes", "acoes"), ("clubes", "clubes"), ("titulos publicos", "tit_publicos"), ("titulo privados", "tit_privados"), ("titulos privados", "tit_privados"), ("cdb", "cdb"), ("dpge", "dpge"), ("letras financeiras", "lf"), ("operacoes compromissadas", "compromissada"),
                  ("outros bancarios", "outros_bancarios"), ("debentures tradicionais", "deb_tradicionais"), ("debentures incentivadas", "deb_incentivadas"), ("debentures", "debentures"), ("cri", "cri"), ("certificado de recebiveis imob", "cri"), ("lci", "lci"), ("letras de credito imob", "lci"),
                  ("letras hipotecarias", "lh"), ("letra imobiliaria garantida", "lig"), ("outros imobiliarios", "outros_imob"), ("lca", "lca"), ("letras de credito agr", "lca"), ("cra", "cra"), ("certificado de recebiveis agr", "cra"), ("outros agricolas", "outros_agro"),
                  ("letra de arrendamento", "lam"), ("letra de cambio", "lc"), ("box", "box"), ("outros titulos privados", "outros_tp"), ("outros ativos", "outros_ativos"), ("coe", "coe"), ("certificado de operacoes", "coe"), ("ativos de captacao bancaria", "captacao_bancaria"), ("ativos com lastro imobiliario", "lastro_imob"), ("ativos com lastro agricola", "lastro_agro"), ("renda variavel", "tvm_rv"), ("ativos de renda fixa", "tvm_rf"), ("renda fixa", "tvm_rf")],
          "caixa_poupanca": [("caixa", "caixa"), ("poupanca", "poupanca")]}

def parse_private(path):
    G = grade(path); out = []
    hdr = next((r for r in G[:10] if sum(1 for v in r if v is not None and ref(v)) >= 2), None)
    tot = next((r for r in G[:12] if any(isinstance(v, str) and norm(v).startswith(("i - volume financeiro", "i - posicao de aum")) for v in r)), None)
    lab_col = next((j for j, v in enumerate(tot) if isinstance(v, str) and v.strip()), 1) if tot else 2
    if hdr:  # layout antigo com duas colunas (dez do ano anterior, mês corrente)
        cols = [(j, ref(v)) for j, v in enumerate(hdr) if v is not None and ref(v)]
    else:
        d = next((ref(v) for r in G[:10] for v in r if v is not None and ref(v)), None)
        vcol = next((j for j, v in enumerate(tot) if num(v) is not None), lab_col + 1) if tot else 3; cols = [(vcol, d)]
    for vcol, data in cols:
        rows = []; sec = None; vistos = set(); got_total = False
        for row in G:
            lab = row[lab_col] if len(row) > lab_col else None
            if not (isinstance(lab, str) and lab.strip()): continue
            n = norm(lab); val = num(row[vcol]) if len(row) > vcol else None
            sh = next((s for k, s in P_SECOES if n.startswith(k)), None)
            if sh == "fim": break
            if sh == "total":
                if not got_total and val is not None: rows.append(("private", "total", val)); got_total = True
                continue
            if sh: sec = sh; rows.append(("private", sh, val)); continue
            if sec in P_MAPA and val is not None:
                for pref, cl in P_MAPA[sec]:
                    if casa(n, pref) and ("private", cl) not in vistos: rows.append(("private", cl, val)); vistos.add(("private", cl)); break
        out.append(pd.DataFrame(rows, columns=["seg", "classe", "valor"]).assign(data=data, arquivo=os.path.basename(path)))
    return pd.concat(out)

def main():
    fv = sorted(glob.glob(os.path.join(AN, "historico", "varejo_*.xls*")) + glob.glob(os.path.join(AN, "historico", "varejo_*.ods")) + glob.glob(os.path.join(AN, "Estatistica_de_Varejo_*.xlsx")))
    fp = sorted(glob.glob(os.path.join(AN, "historico", "private_Estatistica*.xls*")) + glob.glob(os.path.join(AN, "historico", "private_relatorio_*.xls*")) + glob.glob(os.path.join(AN, "Estatistica_de_Private_*.xlsx")))
    L = pd.concat([parse_varejo(f) for f in fv] + [parse_private(f) for f in fp])
    L["pub"] = L.arquivo.str.extract(r"(\d{6})\.(?:xls|ods)")[0].astype(int)
    L = L.sort_values(["seg", "data", "classe", "pub"]).drop_duplicates(["seg", "data", "classe"], keep="last")  # mesma referência em várias edições → a mais recente
    L.drop(columns="pub").to_csv(os.path.join(DATA, "anbima_alocacao_segmentos_long.csv"), index=False)
    W = L.pivot_table(index=["seg", "data"], columns="classe", values="valor").reset_index()
    for c in ["lci", "lca", "cra", "cri", "lf", "debentures", "lig", "lc", "lam", "tvm_outros", "compromissada", "box", "dpge", "outros_bancarios", "outros_imob", "outros_agro", "outros_tp", "outros_ativos",
              "fidc", "fii", "fip", "etf", "fmp", "fundos", "estruturados", "fundos_555", "caixa_poupanca", "previdencia", "outros_invest", "poupanca", "tvm"]:
        if c not in W: W[c] = float("nan")
    for c in ["captacao_bancaria", "lastro_imob", "lastro_agro"]:
        if c not in W: W[c] = float("nan")
    W["lci_lca"] = W.lci.fillna(0) + W.lca.fillna(0)
    old14 = W.lci.isna() & W.lastro_imob.notna(); W.loc[old14, "lci_lca"] = W.loc[old14, "lastro_imob"].fillna(0) + W.loc[old14, "lastro_agro"].fillna(0)  # 2014-15: lastro imob.+agro (inclui CRI/CRA)
    W.loc[old14, "cdb"] = W.loc[old14, "captacao_bancaria"]  # 2014-15: captação bancária (inclui LF, compromissadas, DPGE)
    W["credito_privado"] = W[["cra", "cri", "lf", "debentures", "lig", "lc", "lam", "tvm_outros", "compromissada", "box", "dpge", "outros_bancarios", "outros_imob", "outros_agro", "outros_tp", "outros_ativos"]].fillna(0).sum(axis=1)
    pv = W.seg == "private"
    W.loc[pv, "estruturados"] = W.loc[pv, "estruturados"].fillna(W.loc[pv, ["fidc", "fii", "fip", "etf", "fmp"]].fillna(0).sum(axis=1))
    novo = pv & W.fundos_555.isna() & W["fundos"].notna() & (W.data >= 202405)
    W.loc[pv, "fundos_555"] = W.loc[pv, "fundos"] - W.loc[pv, "estruturados"].fillna(0); W.loc[novo, "fundos_555"] = W.loc[novo, "fundos"]  # novo layout já separa estruturados
    W.loc[~pv, "estruturados"] = W.loc[~pv, "estruturados"].fillna(W.loc[~pv, ["fidc", "fii", "fip", "etf"]].fillna(0).sum(axis=1))
    W.round(1).to_csv(os.path.join(DATA, "anbima_alocacao_segmentos.csv"), index=False)
    chk = W.copy(); chk["soma"] = chk[["fundos_555", "estruturados", "tvm", "poupanca", "previdencia"]].fillna(0).sum(axis=1)
    chk.loc[pv, "soma"] = chk.loc[pv, ["fundos_555", "estruturados", "tvm", "caixa_poupanca", "previdencia", "outros_invest"]].fillna(0).sum(axis=1)
    chk["dif_%"] = (chk.soma / chk.total - 1) * 100
    pd.set_option("display.width", 250); print(chk[["seg", "data", "total", "soma", "dif_%"]].assign(total=lambda d: d.total / 1e3, soma=lambda d: d.soma / 1e3).round(1).to_string(index=False))
    return W

if __name__ == "__main__": main()
